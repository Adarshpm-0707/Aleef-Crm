"""
Report generation service for Aleef CRM.
Builds formatted CSV, Excel (.xlsx via openpyxl), and PDF (via reportlab) exports.
"""

import csv
import io
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class ReportGenerator:
    """Provides methods to export tabular data to CSV, Excel, and PDF."""

    @staticmethod
    def generate_csv(
        data: List[Dict[str, Any]],
        columns: List[Tuple[str, str]],  # [(field_key, display_label), ...]
    ) -> bytes:
        """
        Generates a standard CSV byte stream with UTF-8 BOM encoding.
        """
        output = io.StringIO()
        # Add UTF-8 BOM for Microsoft Excel compatibility
        output.write("\ufeff")

        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

        # Header row
        headers = [col[1] for col in columns]
        writer.writerow(headers)

        # Data rows
        for row in data:
            row_values = []
            for col_key, _ in columns:
                val = row.get(col_key, "")
                if val is None:
                    val = ""
                elif isinstance(val, (datetime,)):
                    val = val.strftime("%Y-%m-%d %H:%M")
                row_values.append(str(val))
            writer.writerow(row_values)

        return output.getvalue().encode("utf-8")

    @staticmethod
    def generate_excel(
        title: str,
        subtitle: str,
        data: List[Dict[str, Any]],
        columns: List[Tuple[str, str]],
        sheet_name: str = "Report Data",
    ) -> bytes:
        """
        Builds a styled, formatted Excel spreadsheet using openpyxl.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel limits sheet names to 31 chars

        # Styles
        title_font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
        title_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")

        meta_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")

        data_font = Font(name="Segoe UI", size=10, color="0F172A")
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        num_cols = max(len(columns), 1)

        # 1. Title Banner
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        title_cell = ws.cell(row=1, column=1, value=title.upper())
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # 2. Metadata / Subtitle
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
        timestamp_str = f"{subtitle} | Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} by Aleef CRM"
        meta_cell = ws.cell(row=2, column=1, value=timestamp_str)
        meta_cell.font = meta_font
        meta_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Empty spacer row
        ws.row_dimensions[3].height = 10

        # 3. Table Column Headers (Row 4)
        for col_idx, (_, col_label) in enumerate(columns, start=1):
            cell = ws.cell(row=4, column=col_idx, value=col_label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[4].height = 26

        # 4. Data Rows
        start_data_row = 5
        for row_idx, row_data in enumerate(data, start=start_data_row):
            is_alt = (row_idx % 2 == 0)
            for col_idx, (col_key, _) in enumerate(columns, start=1):
                val = row_data.get(col_key, "")
                if val is None:
                    val = "-"
                elif isinstance(val, bool):
                    val = "Yes" if val else "No"
                elif isinstance(val, (int, float)):
                    # Keep numeric type for Excel calculations
                    pass
                else:
                    val = str(val)

                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                if is_alt:
                    cell.fill = alt_fill

                # Formatting based on type
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, float):
                        cell.number_format = "#,##0.00"
                    else:
                        cell.number_format = "#,##0"
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 20

        # 5. Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Ignore merged title row for width calculation
                if cell.row in (1, 2, 3):
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_pdf(
        title: str,
        subtitle: str,
        data: List[Dict[str, Any]],
        columns: List[Tuple[str, str]],
        summary_metrics: Optional[Dict[str, Any]] = None,
    ) -> bytes:
        """
        Builds a styled PDF document using ReportLab Platypus elements.
        """
        buffer = io.BytesIO()

        # Choose landscape if more than 5 columns to prevent overflow
        pagesize = landscape(letter) if len(columns) > 5 else letter
        doc = SimpleDocTemplate(
            buffer,
            pagesize=pagesize,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name="ReportTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1E1B4B"),
            spaceAfter=4,
        )

        subtitle_style = ParagraphStyle(
            name="ReportSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=12,
        )

        table_header_style = ParagraphStyle(
            name="TableHeader",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=11,
            textColor=colors.white,
            alignment=1,  # Center
        )

        table_cell_style = ParagraphStyle(
            name="TableCell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#0F172A"),
        )

        metric_label_style = ParagraphStyle(
            name="MetricLabel",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#64748B"),
            alignment=1,
        )

        metric_value_style = ParagraphStyle(
            name="MetricValue",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#1E1B4B"),
            alignment=1,
        )

        story = []

        # 1. Header Banner
        story.append(Paragraph(f"ALEEF CRM — {title.upper()}", title_style))
        gen_time = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
        story.append(Paragraph(f"{subtitle}  |  Generated on {gen_time}", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#4338CA"), spaceAfter=14))

        # 2. Key Metrics Summary Grid (if provided)
        if summary_metrics and len(summary_metrics) > 0:
            metric_headers = []
            metric_values = []
            for label, val in list(summary_metrics.items())[:6]:
                formatted_val = f"{val:,.2f}" if isinstance(val, float) else f"{val:,}" if isinstance(val, int) else str(val)
                metric_headers.append(Paragraph(label.upper(), metric_label_style))
                metric_values.append(Paragraph(formatted_val, metric_value_style))

            metrics_table_data = [metric_headers, metric_values]
            metrics_table = Table(metrics_table_data, hAlign="LEFT")
            metrics_table.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
                    ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ])
            )
            story.append(metrics_table)
            story.append(Spacer(1, 14))

        # 3. Main Data Table
        # Header Row
        table_rows = [[Paragraph(col[1], table_header_style) for col in columns]]

        # Data Rows
        for row in data[:250]:  # Limit PDF page bloat to 250 rows
            row_cells = []
            for col_key, _ in columns:
                raw_val = row.get(col_key, "")
                if raw_val is None:
                    display_text = "-"
                elif isinstance(raw_val, float):
                    display_text = f"{raw_val:,.2f}"
                elif isinstance(raw_val, int):
                    display_text = f"{raw_val:,}"
                elif isinstance(raw_val, bool):
                    display_text = "Yes" if raw_val else "No"
                else:
                    display_text = str(raw_val)
                row_cells.append(Paragraph(display_text, table_cell_style))
            table_rows.append(row_cells)

        if len(table_rows) > 1:
            data_table = Table(table_rows, repeatRows=1, hAlign="LEFT")
            table_style_commands = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4338CA")),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ]

            # Alternating row colors
            for r_idx in range(1, len(table_rows)):
                if r_idx % 2 == 0:
                    table_style_commands.append(
                        ("BACKGROUND", (0, r_idx), (-1, r_idx), colors.HexColor("#F8FAFC"))
                    )

            data_table.setStyle(TableStyle(table_style_commands))
            story.append(data_table)
        else:
            story.append(Paragraph("No records found matching the specified filters.", subtitle_style))

        # Build document
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
